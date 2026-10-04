"""Build the photo gallery (gallery.html) from photo_gallery/.

    photo_gallery/
        albums.json           album titles, descriptions, order; the album marked "banner": true
                              is the big strip at the top; "flip" lists photos to mirror left-right
        photo_details.xlsx    one row per photo: event, date, place, people, note (fill in any time)
        contact_sheet.html    numbered overview of all photos and their details -- local only
        <album folder>/       original photos (any size) -- not published
        web/<album folder>/   web-sized copies made by this script -- published

Add photos to an album folder (or make a new folder), then run:
    python tools/build_gallery.py
New photos get a row in photo_details.xlsx (with the date filled in when the photo or its
file name tells it), and new folders get an album entry in albums.json.
"""
import html
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageOps

ROOT = Path(__file__).resolve().parent.parent
GAL = ROOT / "photo_gallery"
WEB = GAL / "web"
CONFIG = GAL / "albums.json"
DETAILS = GAL / "photo_details.xlsx"
SHEET = GAL / "contact_sheet.html"
OUT = ROOT / "assets" / "js" / "gallery-data.js"
EXTS = {".jpg", ".jpeg", ".png", ".webp", ".heic"}
LONG_SIDE = 1800
COLUMNS = ["Album", "Photo", "Event", "Date", "Place", "People", "Note"]
FIELDS = ["event", "date", "place", "people", "note"]          # the columns you fill in
MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]


def photos_in(folder):
    try:
        return sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in EXTS)
    except OSError:
        return []


def web_name(src):
    """Web-safe file name: lower case, no spaces, no leading "_" or "." (GitHub Pages hides those)."""
    return re.sub(r"[^a-z0-9._-]+", "-", src.stem.lower()).strip("-_.") + ".jpg"


def web_copy(src, album, flip=False):
    """Make the web-sized copy (mirrored left-right if flip) unless an up-to-date one exists."""
    dst = WEB / album.lower() / web_name(src)
    newest = max(src.stat().st_mtime, CONFIG.stat().st_mtime if flip else 0)
    if not dst.exists() or dst.stat().st_mtime < newest:
        dst.parent.mkdir(parents=True, exist_ok=True)
        im = ImageOps.exif_transpose(Image.open(src)).convert("RGB")
        im.thumbnail((LONG_SIDE, LONG_SIDE), Image.LANCZOS)
        if flip:
            im = ImageOps.mirror(im)
        im.save(dst, "JPEG", quality=82, optimize=True, progressive=True)
        print(f"  resized {src.name} -> web/{album.lower()}/{dst.name} ({dst.stat().st_size // 1024} KB)")
    return dst.name


def guess_date(src):
    """'Month YYYY' from the camera data, or from the file name; '' if unknown."""
    try:
        exif = Image.open(src).getexif()
        stamp = exif.get_ifd(0x8769).get(36867) or exif.get(306)      # DateTimeOriginal / DateTime
        if stamp:
            d = datetime.strptime(str(stamp)[:10], "%Y:%m:%d")
            return f"{MONTHS[d.month - 1]} {d.year}"
    except Exception:
        pass
    name = src.name
    if re.match(r"IMG-\d{8}-WA", name):          # WhatsApp save date, not when the photo was taken
        return ""
    m = re.search(r"(20\d\d)[-_]?([01]\d)[-_]?([0-3]\d)", name)
    if m and 1 <= int(m.group(2)) <= 12:
        return f"{MONTHS[int(m.group(2)) - 1]} {m.group(1)}"
    return ""


# ---------------------------------------------------------------- photo details spreadsheet
def read_details():
    if not DETAILS.exists():
        return {}
    ws = load_workbook(DETAILS).active
    rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[1]:
            continue
        rows[str(r[1]).strip()] = {f: ("" if v is None else str(v).strip()) for f, v in zip(FIELDS, r[2:7])}
    return rows


def write_details(entries):
    """entries: list of (album folder, photo name, fields) in page order."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Photo details"
    ws.append(COLUMNS)
    for album, name, d in entries:
        ws.append([album, name] + [d.get(f, "") for f in FIELDS])
        ws.cell(ws.max_row, 2).hyperlink = f"{album}/{name}"       # click to open the photo
        ws.cell(ws.max_row, 2).font = Font(color="0563C1", underline="single")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="6B5EB6")
    for col, w in zip("ABCDEFG", [20, 34, 34, 16, 24, 34, 40]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.number_format = "@"                                    # keep dates as typed text
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    try:
        wb.save(DETAILS)
        return True
    except PermissionError:
        print("!! photo_details.xlsx is open in Excel -- close it and run again.")
        return False


# ---------------------------------------------------------------- contact sheet (local only)
EDITS_NAME = "photo_details_edits.json"
DOWNLOADS = Path.home() / "Downloads"

SHEET_CSS = """
body { font-family: Lato, Arial, sans-serif; margin: 0; color: #333; background: #f6f5fb; }
header { position: sticky; top: 0; z-index: 5; background: #fff; box-shadow: 0 2px 8px rgba(0,0,0,.08); padding: 14px 24px; }
h1 { color: #6B5EB6; margin: 0; font-size: 22px; }
.bar { display: flex; flex-wrap: wrap; align-items: center; gap: 12px; margin-top: 8px; font-size: 14px; }
.bar button { background: #6B5EB6; color: #fff; border: none; border-radius: 16px; padding: 7px 18px; font-size: 14px; cursor: pointer; }
.bar button:disabled { background: #bbb; cursor: default; }
.bar label { color: #555; }
#status { color: #2e7d32; font-weight: bold; }
main { padding: 8px 24px 40px; }
p.help { max-width: 1000px; font-size: 14px; color: #555; }
h2 { color: #6B5EB6; margin: 28px 0 10px; } small { color: #888; font-weight: normal; }
.grid { display: grid; grid-template-columns: repeat(auto-fill, minmax(250px, 1fr)); gap: 14px; }
figure { margin: 0; background: #fff; border-radius: 8px; overflow: hidden; position: relative; box-shadow: 0 1px 4px rgba(0,0,0,.1); }
figure.todo { outline: 2px dashed #d9a6cb; }
figure.edited { outline: 2px solid #6B5EB6; }
figure img { width: 100%; height: 170px; object-fit: cover; display: block; }
.n { position: absolute; top: 6px; left: 6px; background: #6B5EB6; color: #fff; font-weight: bold; padding: 2px 8px; border-radius: 10px; font-size: 14px; }
figcaption { padding: 8px 10px 10px; font-size: 13px; }
.f { color: #999; font-size: 11px; word-break: break-all; margin-bottom: 5px; }
figcaption input { width: 100%; box-sizing: border-box; margin: 2px 0; padding: 4px 6px; border: 1px solid #ddd; border-radius: 4px; font: inherit; }
figcaption input:focus { outline: none; border-color: #6B5EB6; }
body.hide-done figure:not(.todo):not(.edited) { display: none; }
"""

SHEET_JS = r"""
var KEY = 'gallery-photo-drafts';
var drafts = {};
try { drafts = JSON.parse(localStorage.getItem(KEY) || '{}'); } catch (e) {}
var FIELDS = ['event', 'date', 'place', 'people', 'note'];
var cards = document.querySelectorAll('figure[data-photo]');

function orig(card, f) { return card.querySelector('[name=' + f + ']').dataset.orig; }
function changed(card) {
    return FIELDS.some(function (f) { return card.querySelector('[name=' + f + ']').value.trim() !== orig(card, f); });
}
function refresh() {
    var n = 0;
    cards.forEach(function (c) { var ch = changed(c); c.classList.toggle('edited', ch); if (ch) n++; });
    document.getElementById('count').textContent = n ? n + ' photo' + (n > 1 ? 's' : '') + ' with changes' : 'no changes yet';
    document.getElementById('save').disabled = !n;
}
cards.forEach(function (card) {
    var name = card.dataset.photo, d = drafts[name];
    if (d) {   // restore what was typed earlier (unless it has since made it into the spreadsheet)
        FIELDS.forEach(function (f) { if (d[f] !== undefined) card.querySelector('[name=' + f + ']').value = d[f]; });
        if (!changed(card)) delete drafts[name];
    }
    card.addEventListener('input', function () {
        var rec = {};
        FIELDS.forEach(function (f) { rec[f] = card.querySelector('[name=' + f + ']').value.trim(); });
        if (changed(card)) drafts[name] = rec; else delete drafts[name];
        try { localStorage.setItem(KEY, JSON.stringify(drafts)); } catch (e) {}
        document.getElementById('status').textContent = '';
        refresh();
    });
});
try { localStorage.setItem(KEY, JSON.stringify(drafts)); } catch (e) {}
refresh();

document.getElementById('hide').onchange = function () { document.body.classList.toggle('hide-done', this.checked); };

document.getElementById('save').onclick = async function () {
    var out = {};
    cards.forEach(function (card) {
        if (!changed(card)) return;
        var rec = {};
        FIELDS.forEach(function (f) { rec[f] = card.querySelector('[name=' + f + ']').value.trim(); });
        out[card.dataset.photo] = rec;
    });
    var text = JSON.stringify(out, null, 2);
    var n = Object.keys(out).length;
    try {
        if (window.showSaveFilePicker) {
            var h = await window.showSaveFilePicker({ suggestedName: 'photo_details_edits.json',
                types: [{ description: 'Photo details', accept: { 'application/json': ['.json'] } }] });
            var w = await h.createWritable(); await w.write(text); await w.close();
        } else { throw 'no picker'; }
    } catch (e) {
        if (e && e.name === 'AbortError') return;
        var a = document.createElement('a');
        a.href = URL.createObjectURL(new Blob([text], { type: 'application/json' }));
        a.download = 'photo_details_edits.json'; a.click();
    }
    document.getElementById('status').textContent = 'Saved ' + n + ' photo' + (n > 1 ? 's' : '') +
        ' \u2014 tell Claude "I added photo details" to put them on the website.';
};
"""


def write_contact_sheet(albums_out):
    parts = []
    for a in albums_out:
        cards = []
        for i, p in enumerate(a["photos"], 1):
            d = p["details"]
            todo = "" if d.get("event") else " todo"
            inputs = "".join(
                f'<input name="{k}" placeholder="{ph}" value="{html.escape(d.get(k, ""), quote=True)}" '
                f'data-orig="{html.escape(d.get(k, ""), quote=True)}">'
                for k, ph in [("event", "Event (e.g. Lab hike, ISNA 2025)"), ("date", "Date (e.g. July 2025)"),
                              ("place", "Place"), ("people", "People (e.g. Itay, Shany)"), ("note", "Note")])
            cards.append(f'<figure class="{todo.strip()}" data-photo="{html.escape(p["name"], quote=True)}">'
                         f'<span class="n">{a["code"]} {i}</span>'
                         f'<a href="{p["web"]}" target="_blank"><img src="{p["web"]}" loading="lazy"></a>'
                         f'<figcaption><div class="f">{html.escape(p["name"])}</div>{inputs}</figcaption></figure>')
        parts.append(f'<h2>{html.escape(a["title"])} <small>({a["code"]} 1&ndash;{len(a["photos"])})</small></h2>'
                     f'<div class="grid">{"".join(cards)}</div>')
    SHEET.write_text(f"""<!doctype html><html><head><meta charset="utf-8"><title>Gallery photo details</title>
<style>{SHEET_CSS}</style></head><body>
<header><h1>Gallery photo details</h1>
<div class="bar"><button id="save" disabled>Save my changes</button><span id="count"></span><span id="status"></span>
<label><input type="checkbox" id="hide"> show only photos that still need details</label></div></header>
<main><p class="help">Type into the boxes under any photo &mdash; leave blank what you don't know. Your typing is kept in
this browser automatically, so you can stop and come back later. When you're ready, click <b>Save my changes</b>
(save the file in the <b>photo_gallery</b> folder or in Downloads) and tell Claude <i>"I added photo details"</i>.
Dashed frames have no event yet; purple frames have unsaved changes. This page is only on your computer.</p>
{''.join(parts)}</main><script>{SHEET_JS}</script></body></html>""", encoding="utf8")


def read_edits():
    """Edits saved from the contact sheet (photo_gallery/ or Downloads), oldest first."""
    found = [GAL / EDITS_NAME] + sorted(DOWNLOADS.glob("photo_details_edits*.json"), key=lambda f: f.stat().st_mtime)
    files = [f for f in found if f.exists()]
    edits = {}
    for f in files:
        try:
            edits.update(json.loads(f.read_text(encoding="utf8")))
            print(f"Applying {len(json.loads(f.read_text(encoding='utf8')))} photo edits from {f}")
        except ValueError:
            print(f"!! could not read {f}")
    return edits, files


def main():
    cfg = json.loads(CONFIG.read_text(encoding="utf8"))
    albums = cfg.setdefault("albums", [])

    def folders_of(a):  # an album can draw on one folder ("folder") or several ("folders")
        return a.get("folders") or [a["folder"]]

    known = {f for a in albums for f in folders_of(a)}
    for folder in sorted(p for p in GAL.iterdir() if p.is_dir() and p.name != "web"):
        if folder.name not in known and photos_in(folder):
            albums.append({"folder": folder.name, "title": folder.name.replace("_", " ").replace("-", " ").title(),
                           "date": "", "description": ""})
            print(f"New album '{folder.name}' added to albums.json -- fill in its title and description.")
    CONFIG.write_text(json.dumps(cfg, indent=4, ensure_ascii=False) + "\n", encoding="utf8")

    details = read_details()
    edits, edit_files = read_edits()
    for name, rec in edits.items():                       # edits from the contact sheet win
        details[name] = {k: str(rec.get(k, details.get(name, {}).get(k, ""))).strip() for k in FIELDS}
    entries, data, sheet = [], [], []
    for a in albums:
        files = [(f, p) for f in folders_of(a) for p in photos_in(GAL / f)]
        if not files:
            print(f"Skipping '{a['title']}': no photos yet.")
            continue
        order = a.get("order", [])
        files.sort(key=lambda fp: (order.index(fp[1].name) if fp[1].name in order else len(order), fp[0], fp[1].name))
        print(f"{a['title']}: {len(files)} photos")
        flips = set(a.get("flip", []))
        photos, sheet_photos = [], []
        for f, p in files:
            d = details.get(p.name)
            if d is None:                                    # new photo: start its row
                d = {k: "" for k in FIELDS}
                d["date"] = guess_date(p)
            entries.append((f, p.name, d))
            web = f"photo_gallery/web/{f.lower()}/{web_copy(p, f, p.name in flips)}"
            photos.append(dict({"file": web}, **{k: d[k] for k in FIELDS if d.get(k)}))
            sheet_photos.append({"name": p.name, "web": web[len("photo_gallery/"):], "details": d})
        data.append({"title": a["title"], "date": a.get("date", ""), "description": a.get("description", ""),
                     "banner": bool(a.get("banner")), "folder": "", "photos": photos})
        code = a.get("code") or a["title"].split()[-1].strip("!")
        sheet.append({"title": a["title"], "code": code, "photos": sheet_photos})

    if write_details(entries):
        for f in edit_files:                              # now safely in the spreadsheet
            f.unlink()
    write_contact_sheet(sheet)

    # remove web copies whose original is gone (and folders left empty)
    # web folder names are lower case (web servers are case-sensitive, Windows is not)
    originals = {p.name.lower(): p for p in GAL.iterdir() if p.is_dir() and p.name != "web"}
    for folder in list(WEB.iterdir()) if WEB.exists() else []:
        src = originals.get(folder.name.lower())
        keep = {web_name(p) for p in photos_in(src)} if src else set()
        for f in folder.iterdir():
            if f.name not in keep:
                f.unlink()
                print(f"  removed stale web/{folder.name}/{f.name}")
        if not any(folder.iterdir()):
            folder.rmdir()

    OUT.write_text("// Generated by tools/build_gallery.py from photo_gallery/ -- edit albums.json and "
                   "photo_details.xlsx, not this file.\n"
                   "window.GALLERY_ALBUMS = " + json.dumps(data, indent=4, ensure_ascii=False) + ";\n", encoding="utf8")
    print(f"Wrote {OUT.relative_to(ROOT)}, photo_details.xlsx and contact_sheet.html")

    # bump the version stamp on the data file in gallery.html so browsers don't show a cached copy
    page = ROOT / "gallery.html"
    text = page.read_text(encoding="utf8")
    stamp = str(int(OUT.stat().st_mtime))
    text = re.sub(r'assets/js/gallery-data\.js(\?v=\d+)?"', f'assets/js/gallery-data.js?v={stamp}"', text)
    page.write_text(text, encoding="utf8")


if __name__ == "__main__":
    main()
